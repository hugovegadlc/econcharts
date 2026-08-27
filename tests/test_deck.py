"""Deck assembly, and the two kinds of slide.

`chrome` is a property of the SIZE preset, not a spec key: the same YAML gives a
chart that carries its own title (the four original presets) or a bare chart on
a captioned BBVA panel (the 16:9 presets). These tests pin that split, because
nothing else does — the chart itself looks fine either way, and only the deck
shows which one you got.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib
import pytest

matplotlib.use("Agg")

from econcharts.deck import DeckItem, build_deck          # noqa: E402
from econcharts.render import render, save                 # noqa: E402
from econcharts.spec import Spec                           # noqa: E402
from econcharts.theme import load_theme                    # noqa: E402

pptx = pytest.importorskip("pptx")

_PERIODS = [f"2024M{m:02d}" for m in range(1, 13)]


def _spec(**kw):
    return Spec(period=f"{_PERIODS[0]}:{_PERIODS[-1]}",
                series=[{"name": "A", "type": "line",
                         "data": {p: 3 + i * 0.4 for i, p in enumerate(_PERIODS)}}], **kw)


def _png(tmp_path, name, size, **kw):
    fig = render(_spec(**kw), size=size)
    out = tmp_path / f"{name}.png"
    save(fig, out)
    return out


def _texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def _kinds(slide):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    return [sh.shape_type for sh in slide.shapes]


# --- the chart itself --------------------------------------------------------

def test_a_slide_chrome_preset_renders_the_chart_bare():
    """Its words go on the slide instead, so the figure must not repeat them."""
    fig = render(_spec(title="T", subtitle="sub", source="BCRP"), size="slides16_9_half")
    ax = fig.axes[0]
    assert [t for t in (ax.get_title(loc=l) for l in ("left", "center", "right")) if t] == []
    assert [t.get_text() for t in fig.texts] == []


def test_a_chart_chrome_preset_still_carries_its_own_words():
    fig = render(_spec(title="T", subtitle="sub", source="BCRP"), size="slides_half")
    ax = fig.axes[0]
    words = [t for t in (ax.get_title(loc=l) for l in ("left", "center", "right")) if t]
    words += [t.get_text() for t in fig.texts]
    assert "T" in words
    assert any("BCRP" in w for w in words)


# --- the slide ---------------------------------------------------------------

def test_slide_chrome_draws_the_caption_the_rule_and_the_source(tmp_path):
    img = _png(tmp_path, "c", "slides16_9_half", title="Empleo", subtitle="var. %",
               source="INEI")
    deck = build_deck([DeckItem(image=img, size_mm=(100, 80), size_name="slides16_9_half",
                                title="Empleo", subtitle="var. %", source="INEI")],
                      tmp_path / "d.pptx", theme="bbva")
    slide = pptx.Presentation(str(deck)).slides[0]
    blob = "\n".join(_texts(slide))
    assert "EMPLEO" in blob                      # caption, upper-cased by the theme
    assert "(VAR. %)" in blob                    # units line, parenthesised
    assert "Fuente: INEI" in blob                # source footnote
    assert "XXX" in blob                         # title placeholder
    assert "p. 1" in blob.replace("p.  1", "p. 1")
    assert any(sh.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE
               for sh in slide.shapes)


def test_chart_chrome_puts_only_pictures_on_the_slide(tmp_path):
    img = _png(tmp_path, "c", "slides_half", title="Empleo", source="INEI")
    deck = build_deck([DeckItem(image=img, size_mm=(85, 70), size_name="slides_half",
                                title="Empleo", source="INEI")],
                      tmp_path / "d.pptx", theme="bbva")
    slide = pptx.Presentation(str(deck)).slides[0]
    assert _texts(slide) == []                   # the figure already says it all
    assert len(slide.shapes) == 1


def test_the_slide_size_comes_from_the_theme(tmp_path):
    img = _png(tmp_path, "c", "slides16_9_half", title="T")
    deck = build_deck([DeckItem(image=img, size_mm=(100, 80), size_name="slides16_9_half",
                                title="T")], tmp_path / "d.pptx", theme="bbva")
    prs = pptx.Presentation(str(deck))
    theme = load_theme("bbva")
    from pptx.util import Inches
    assert prs.slide_width == Inches(float(theme.val("deck.slide_width_in")))
    assert prs.slide_height == Inches(float(theme.val("deck.slide_height_in")))


def test_per_slide_comes_from_the_theme(tmp_path):
    """`slides16_9_full` is one chart per slide, `_half` is two."""
    def deck_for(size, n):
        img = _png(tmp_path, f"c{size}", size, title="T")
        w = 218 if size.endswith("full") else 100
        items = [DeckItem(image=img, size_mm=(w, 80), size_name=size, title="T")] * n
        return pptx.Presentation(str(build_deck(items, tmp_path / f"{size}.pptx", theme="bbva")))

    assert len(deck_for("slides16_9_half", 4).slides) == 2
    assert len(deck_for("slides16_9_full", 4).slides) == 4


def test_a_long_caption_shrinks_rather_than_running_through_the_rule(tmp_path):
    """Bounded: it steps down to `caption.min_size` and stops. The rule sits at a
    fixed distance below the panel top — that spacing is the house style and does
    not move — so the caption is what gives."""
    long_title = "Un titulo deliberadamente largo " * 3
    img = _png(tmp_path, "c", "slides16_9_half", title=long_title)
    deck = build_deck([DeckItem(image=img, size_mm=(100, 80), size_name="slides16_9_half",
                                title=long_title, subtitle="var. % interanual")],
                      tmp_path / "d.pptx", theme="bbva")
    slide = pptx.Presentation(str(deck)).slides[0]
    theme = load_theme("bbva")
    floor = float(theme.val("deck.slide.caption.min_size"))
    full = float(theme.val("deck.slide.caption.size"))
    caption = next(sh for sh in slide.shapes
                   if sh.has_text_frame and long_title.upper()[:20] in sh.text_frame.text)
    pt = caption.text_frame.paragraphs[0].runs[0].font.size.pt
    assert floor <= pt < full


def test_tuples_are_still_accepted(tmp_path):
    """The signature before slide chrome existed."""
    img = _png(tmp_path, "c", "slides_half", title="T")
    deck = build_deck([(img, (85, 70))], tmp_path / "d.pptx")
    assert len(pptx.Presentation(str(deck)).slides) == 1
